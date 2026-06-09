import sys
import os
import subprocess
import json
import datetime
import csv
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QTreeWidget, QTreeWidgetItem, QScrollArea,
    QLineEdit, QMessageBox, QMenu, QButtonGroup, QInputDialog, QProgressBar,
    QDialog, QFrame, QSplitter, QTabWidget, QTextEdit, QGraphicsOpacityEffect,
    QSystemTrayIcon, QSizePolicy, QStyledItemDelegate, QStyle, QFileDialog,
    QSplashScreen
)
from PySide6.QtCore import (
    Qt, QThread, Signal, QMimeData, QTimer, QPropertyAnimation,
    QEasingCurve, QRect, QPoint, QSize, QParallelAnimationGroup
)
from PySide6.QtGui import (
    QColor, QPixmap, QIcon, QGuiApplication, QFont, QPalette,
    QLinearGradient, QBrush, QPainter, QAction, QKeySequence, QShortcut
)

try:
    import winsound
    HAS_WINSOUND = True
except ImportError:
    HAS_WINSOUND = False

# ─────────────────────── CONFIG ───────────────────────
EXCEL_PATH = r""
VNC_PASSWORD = ""
OTOMATIK_IZLEME_DAKIKA = 5
LOG_FILE = "vnc_log.json"
MAX_LOG_ENTRIES = 500

SHEET_MAP = {
    "Montaj 1": "10.105.x.x",
    "Montaj 2": "10.105.x.x",
    "Montaj 3": "10.105.x.x",
    "Üretim 39. Blok": "10.105.x.x"
}
OCTET_MAP = {
    "Montaj 1": "36",
    "Montaj 2": "37",
    "Montaj 3": "38",
    "Üretim 39. Blok": "39"
}

# ─────────────────── RENK PALETİ ───────────────────
COLORS = {
    "light": {
        "bg_sidebar":    "#f0f2f8", "bg_sidebar2":   "#e4e8f4", "bg_main":       "#ffffff",
        "bg_card":       "#f8f9fc", "accent":        "#3b7cf8", "accent2":       "#22c77a",
        "text":          "#1a1d2e", "text_dim":      "#6b7280", "border":        "#dde1ec",
        "online":        "#22c77a", "offline":       "#ef4444", "unknown":       "#9ca3af",
        "hover":         "#e4e8f4", "btn_bg":        "#eef1f8", "btn_hover":     "#dde3f5",
        "header_bg":     "#f0f2f8", "logo_bg":       "transparent"
    },
    "dark": {
        "bg_sidebar":    "#1e1e2e", "bg_sidebar2":   "#2a2a3c", "bg_main":       "#141421",
        "bg_card":       "#222234", "accent":        "#5c93fa", "accent2":       "#2dd486",
        "text":          "#e2e4eb", "text_dim":      "#a0a5b5", "border":        "#33334d",
        "online":        "#2dd486", "offline":       "#fb5c5c", "unknown":       "#6b7280",
        "hover":         "#2a2a3c", "btn_bg":        "#2a2a3c", "btn_hover":     "#3b3b54",
        "header_bg":     "#1e1e2e", "logo_bg":       "transparent"
    }
}

# ─────────────────── YARDIMCI ───────────────────
def get_data_path(filename):
    base = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(__file__)
    return os.path.join(base, filename)

def get_config_file():
    base = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(__file__)
    return os.path.join(base, "config.json")

def get_user_folder_file():
    base = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(__file__)
    target = os.path.join(base, "user_folders.json")
    if not os.path.exists(target):
        with open(target, "w", encoding="utf-8") as f: json.dump({}, f)
    return target

def get_log_file():
    base = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(__file__)
    return os.path.join(base, LOG_FILE)

def get_favorites_file():
    base = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(__file__)
    target = os.path.join(base, "favorites.json")
    if not os.path.exists(target):
        with open(target, "w", encoding="utf-8") as f: json.dump([], f)
    return target

USER_FOLDER_FILE = get_user_folder_file()
CONFIG_FILE = get_config_file()
FAVORITES_FILE = get_favorites_file()

# ─────────────────── LOG MANAGER ───────────────────
class LogManager:
    def __init__(self):
        self.log_file = get_log_file()
        self.entries = self._load()

    def _load(self):
        if os.path.exists(self.log_file):
            try:
                with open(self.log_file, "r", encoding="utf-8") as f: return json.load(f)
            except: return []
        return []

    def add(self, action, target, detail="", status="ok"):
        entry = {
            "time": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "action": action, "target": target, "detail": detail, "status": status
        }
        self.entries.insert(0, entry)
        if len(self.entries) > MAX_LOG_ENTRIES: self.entries = self.entries[:MAX_LOG_ENTRIES]
        self._save()
        return entry

    def _save(self):
        try:
            with open(self.log_file, "w", encoding="utf-8") as f:
                json.dump(self.entries, f, ensure_ascii=False, indent=2)
        except: pass

    def clear(self):
        self.entries = []; self._save()

# ─────────────────── CUSTOM WIDGETS ───────────────────
class AnimatedButton(QPushButton):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.setFixedHeight(38); self.setCursor(Qt.PointingHandCursor)

class SidebarButton(QPushButton):
    def __init__(self, text, icon_char="", *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.icon_char = icon_char
        self.setText(f"  {icon_char}  {text}" if icon_char else f"     {text}")
        self.setCheckable(True); self.setFixedHeight(42); self.setCursor(Qt.PointingHandCursor)

class StatusDelegate(QStyledItemDelegate):
    STATUS_COLORS = {"online": "#22c77a", "offline": "#ef4444", "unknown": "#6b7280"}

    def paint(self, painter, option, index):
        painter.save()
        is_selected = bool(option.state & QStyle.State_Selected)
        is_hover    = bool(option.state & QStyle.State_MouseOver)
        palette_bg = option.palette.base().color().name()
        is_dark    = palette_bg < "#888888"

        bg_main  = QColor("#141421") if is_dark else QColor("#ffffff")
        bg_alt   = QColor("#222234") if is_dark else QColor("#f8f9fc")
        bg_hover = QColor("#2a2a3c") if is_dark else QColor("#e4e8f4")
        text_norm = QColor("#e2e4eb") if is_dark else QColor("#1a1d2e")

        if is_selected:
            bg = QColor("#1a56db")
        elif is_hover:
            bg = bg_hover
        else:
            bg = bg_main if index.row() % 2 == 0 else bg_alt

        painter.fillRect(option.rect, bg)

        if is_selected:
            text_color = QColor("#ffffff")
        elif index.column() == 0:
            status = index.data(Qt.UserRole) or "unknown"
            text_color = QColor(self.STATUS_COLORS.get(status, text_norm.name()))
        elif index.column() == 3:
            text_color = QColor("#f59e0b")  # Offline süre — turuncu
        else:
            text_color = text_norm

        painter.setPen(text_color)
        text = index.data(Qt.DisplayRole) or ""
        painter.drawText(option.rect.adjusted(6, 0, -6, 0), Qt.AlignVCenter | Qt.AlignLeft, text)
        painter.restore()

    def sizeHint(self, option, index):
        hint = super().sizeHint(option, index)
        hint.setHeight(max(hint.height(), 32))
        return hint

class DragTree(QTreeWidget):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.setDragEnabled(True); self.setDragDropMode(QTreeWidget.DragOnly)
        self.setAlternatingRowColors(True); self.setAnimated(True)
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)

    def mimeData(self, items):
        mime = QMimeData()
        lines = ["\t".join([it.text(c) for c in range(it.columnCount())]) for it in items]
        mime.setText("\n".join(lines))
        return mime

class DropButton(AnimatedButton):
    def __init__(self, title, icon_char="📁", drop_callback=None, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.icon_char = icon_char; self._base_title = title
        self.setText(f"  {icon_char}  {title}"); self.setAcceptDrops(True)
        self.drop_callback = drop_callback; self.setCheckable(True)
        self.setCursor(Qt.PointingHandCursor); self.setFixedHeight(42)

    def dragEnterEvent(self, event):
        if event.mimeData().hasText(): event.acceptProposedAction()
        else: event.ignore()

    def dragMoveEvent(self, event):
        if event.mimeData().hasText(): event.acceptProposedAction()
        else: event.ignore()

    def dropEvent(self, event):
        if not event.mimeData().hasText(): event.ignore(); return
        text = event.mimeData().text().splitlines()[0]
        parts = text.split("\t")
        if self.drop_callback:
            self.drop_callback(
                self._base_title,
                parts[0] if len(parts) > 0 else "",
                parts[1] if len(parts) > 1 else "",
                parts[2] if len(parts) > 2 else ""
            )
        event.acceptProposedAction()

# ─────────────────── TOAST BİLDİRİM ───────────────────
class ToastNotification(QWidget):
    def __init__(self, message, notif_type="info", parent=None):
        super().__init__(parent, Qt.FramelessWindowHint | Qt.WindowStaysOnTopHint | Qt.Tool)
        self.setAttribute(Qt.WA_TranslucentBackground); self.setAttribute(Qt.WA_ShowWithoutActivating)
        colors = {"info": ("#4f8ef7", "white"), "success": ("#3ecf8e", "white"),
                  "warning": ("#f59e0b", "white"), "error": ("#f05252", "white")}
        bg, fg = colors.get(notif_type, colors["info"])
        layout = QHBoxLayout(self); layout.setContentsMargins(16, 10, 16, 10)
        icons = {"info": "ℹ️", "success": "✅", "warning": "⚠️", "error": "❌"}
        icon_lbl = QLabel(icons.get(notif_type, "ℹ️"))
        icon_lbl.setStyleSheet("font-size:16px; background: transparent;")
        msg_lbl = QLabel(message)
        msg_lbl.setStyleSheet(f"color:{fg};font-size:13px;font-weight:500; background: transparent;")
        layout.addWidget(icon_lbl); layout.addWidget(msg_lbl)
        self.setStyleSheet(f"QWidget {{ background: {bg}; border-radius: 10px; }}")
        self.adjustSize()
        self.effect = QGraphicsOpacityEffect(self); self.setGraphicsEffect(self.effect)
        self.fade_in = QPropertyAnimation(self.effect, b"opacity")
        self.fade_in.setDuration(300); self.fade_in.setStartValue(0.0); self.fade_in.setEndValue(1.0)
        self.fade_out = QPropertyAnimation(self.effect, b"opacity")
        self.fade_out.setDuration(500); self.fade_out.setStartValue(1.0); self.fade_out.setEndValue(0.0)
        self.fade_out.finished.connect(self.close)
        QTimer.singleShot(3500, self.fade_out.start)

    def show_at(self, parent_widget):
        self.show(); self.fade_in.start()
        if parent_widget:
            pr = parent_widget.geometry()
            self.move(pr.right() - self.width() - 20, pr.bottom() - self.height() - 20)

# ─────────────────── DASHBOARD ───────────────────
class DashboardWidget(QWidget):
    def __init__(self, theme="light", parent=None):
        super().__init__(parent)
        self.theme = theme
        self.stats = {"total": 0, "online": 0, "offline": 0, "unknown": 0, "connections_today": 0}
        self._build_ui()

    def _build_ui(self):
        layout = QVBoxLayout(self); layout.setContentsMargins(20, 20, 20, 20); layout.setSpacing(16)
        title = QLabel("📊  Dashboard"); title.setStyleSheet("font-size:18px;font-weight:bold;")
        layout.addWidget(title)
        cards_row = QHBoxLayout(); cards_row.setSpacing(12)
        C = COLORS[self.theme]; self.card_widgets = {}
        card_data = [
            ("Toplam Cihaz", "total", C["accent"]),
            ("Online", "online", C["online"]),
            ("Offline", "offline", C["offline"]),
            ("Bağlantı (bugün)", "connections_today", C["accent2"])
        ]
        for title_str, key, color in card_data:
            card = QFrame(); card.setObjectName("StatCard"); card.setMinimumHeight(90)
            clay = QVBoxLayout(card); clay.setContentsMargins(16, 12, 16, 12)
            val_lbl = QLabel("0")
            val_lbl.setStyleSheet(f"font-size:30px;font-weight:bold;color:{color};")
            val_lbl.setAlignment(Qt.AlignCenter)
            ttl_lbl = QLabel(title_str)
            ttl_lbl.setStyleSheet("font-size:11px;color:#6b7280;")
            ttl_lbl.setAlignment(Qt.AlignCenter)
            clay.addWidget(val_lbl); clay.addWidget(ttl_lbl)
            self.card_widgets[key] = val_lbl; cards_row.addWidget(card)
        layout.addLayout(cards_row)
        bar_frame = QFrame(); bar_frame.setObjectName("StatCard")
        bar_lay = QVBoxLayout(bar_frame); bar_lay.setContentsMargins(16, 12, 16, 12)
        bar_title = QLabel("Cihaz Durumu Dağılımı")
        bar_title.setStyleSheet("font-size:13px;font-weight:600;margin-bottom:8px;")
        bar_lay.addWidget(bar_title)
        self.online_bar = QProgressBar(); self.online_bar.setTextVisible(False); self.online_bar.setFixedHeight(12)
        self.online_bar.setStyleSheet(
            f"QProgressBar {{ border-radius:6px; background:#2d3139; }}"
            f"QProgressBar::chunk {{ border-radius:6px; background: qlineargradient("
            f"x1:0,y1:0,x2:1,y2:0,stop:0 {COLORS[self.theme]['online']}, stop:1 {COLORS[self.theme]['accent2']}); }}"
        )
        bar_lay.addWidget(self.online_bar)
        self.bar_label = QLabel("Henüz veri yok")
        self.bar_label.setStyleSheet("font-size:11px;color:#6b7280;margin-top:4px;")
        bar_lay.addWidget(self.bar_label)
        layout.addWidget(bar_frame); layout.addStretch()

    def update_stats(self, stats: dict):
        self.stats.update(stats)
        for key, lbl in self.card_widgets.items(): lbl.setText(str(self.stats.get(key, 0)))
        total = self.stats.get("total", 0); online = self.stats.get("online", 0)
        if total > 0:
            self.online_bar.setValue(int(online / total * 100))
            self.bar_label.setText(f"Online: {online} / {total}  ({int(online/total*100)}%)")
        else:
            self.online_bar.setValue(0); self.bar_label.setText("Henüz veri yok")

    def apply_theme(self, theme):
        self.theme = theme; C = COLORS[theme]
        self.setStyleSheet(
            f"DashboardWidget {{ background: {C['bg_main']}; }}"
            f"QFrame#StatCard {{ background: {C['bg_card']}; border: 1px solid {C['border']}; border-radius: 12px; }}"
            f"QLabel {{ color: {C['text']}; }}"
        )

# ─────────────────── LOG PANEL ───────────────────
class LogPanel(QWidget):
    def __init__(self, log_manager: LogManager, parent_vnc, theme="light", parent=None):
        super().__init__(parent)
        self.log_manager = log_manager; self.theme = theme; self.parent_vnc = parent_vnc
        self._build_ui(); self.refresh()

    def _build_ui(self):
        lay = QVBoxLayout(self); lay.setContentsMargins(20, 20, 20, 20); lay.setSpacing(10)
        hdr = QHBoxLayout()
        title = QLabel("📋  Bağlantı Geçmişi"); title.setStyleSheet("font-size:18px;font-weight:bold;")
        hdr.addWidget(title); hdr.addStretch()
        exp_btn = QPushButton("💾 Dışa Aktar"); exp_btn.setFixedHeight(30)
        exp_btn.setCursor(Qt.PointingHandCursor); exp_btn.clicked.connect(self._export_logs)
        hdr.addWidget(exp_btn)
        clr_btn = QPushButton("🗑️ Temizle"); clr_btn.setFixedHeight(30)
        clr_btn.setCursor(Qt.PointingHandCursor); clr_btn.clicked.connect(self._clear_log)
        hdr.addWidget(clr_btn); lay.addLayout(hdr)
        self.search = QLineEdit(); self.search.setPlaceholderText("Logda ara...")
        self.search.textChanged.connect(self.refresh); lay.addWidget(self.search)
        self.log_tree = QTreeWidget()
        self.log_tree.setHeaderLabels(["Zaman", "İşlem", "Hedef", "Detay", "Durum"])
        self.log_tree.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.log_tree.setColumnWidth(0, 140); self.log_tree.setColumnWidth(1, 100)
        self.log_tree.setColumnWidth(2, 130); self.log_tree.setColumnWidth(3, 200)
        self.log_tree.setColumnWidth(4, 70)
        self.log_tree.setAlternatingRowColors(True); self.log_tree.setSortingEnabled(True)
        lay.addWidget(self.log_tree)

    def refresh(self, filter_text=""):
        filter_text = self.search.text().lower() if hasattr(self, "search") else filter_text.lower()
        self.log_tree.clear()
        for e in self.log_manager.entries:
            row = [e.get("time",""), e.get("action",""), e.get("target",""), e.get("detail",""), e.get("status","")]
            if filter_text and not any(filter_text in str(v).lower() for v in row): continue
            item = QTreeWidgetItem(row)
            st = e.get("status", "ok")
            item.setForeground(4, QColor("#3ecf8e") if st == "ok" else QColor("#f05252"))
            self.log_tree.addTopLevelItem(item)

    def _clear_log(self):
        if QMessageBox.question(self, "Onay", "Tüm log silinsin mi?",
                                QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
            self.log_manager.clear(); self.refresh()

    def _export_logs(self):
        path, _ = QFileDialog.getSaveFileName(self, "Logları Dışa Aktar", "VNC_Logs.csv", "CSV Files (*.csv)")
        if not path: return
        try:
            with open(path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["Zaman", "İşlem", "Hedef", "Detay", "Durum"])
                for e in self.log_manager.entries:
                    writer.writerow([e.get("time",""), e.get("action",""),
                                     e.get("target",""), e.get("detail",""), e.get("status","")])
            self.parent_vnc.show_toast("Loglar CSV olarak kaydedildi.", "success")
        except Exception as ex:
            self.parent_vnc.show_toast(f"Dışa aktarma hatası: {ex}", "error")

    def apply_theme(self, theme):
        self.theme = theme; C = COLORS[theme]
        self.log_tree.setStyleSheet(
            f"QTreeWidget {{ background:{C['bg_card']}; color:{C['text']}; border:1px solid {C['border']}; border-radius:8px; }}"
            f"QTreeWidget::item {{ min-height:26px; padding:3px; border-bottom:1px solid {C['border']}; }}"
            f"QTreeWidget::item:selected {{ background:{C['accent']}; color:white; }}"
            f"QHeaderView::section {{ background:{C['header_bg']}; color:{C['text']}; padding:6px; font-weight:bold; border:none; border-bottom:1px solid {C['border']}; }}"
            f"QTreeWidget::item:alternate {{ background:{C['bg_main']}; }}"
        )

# ─────────────────── WORKERS ───────────────────
class LoadSheetWorker(QThread):
    data_loaded = Signal(list, str)
    error = Signal(str)

    def __init__(self, sheet_name, octet, target_group, min_last_octet=135):
        super().__init__()
        self.sheet_name = sheet_name
        self.octet = octet
        self.target_group = target_group      # hangi grup için yüklendiği
        self.min_last_octet = min_last_octet

    def run(self):
        rows = []
        if not os.path.exists(EXCEL_PATH):
            self.error.emit(f"Excel dosyası bulunamadı: {EXCEL_PATH}")
            self.data_loaded.emit(rows, self.target_group); return
        try:
            import openpyxl
            wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
            if self.sheet_name not in wb.sheetnames:
                self.data_loaded.emit(rows, self.target_group); wb.close(); return
            ws = wb[self.sheet_name]
            for row in ws.iter_rows(values_only=True):
                if not row or len(row) < 1: continue
                raw_ip = row[0]; raw_asset = row[1] if len(row) > 1 else None
                raw_desc = row[2] if len(row) > 2 else None
                ip = str(raw_ip).strip() if raw_ip else ""
                asset_str = str(raw_asset) if raw_asset is not None else ""
                desc_str = str(raw_desc) if raw_desc is not None else ""
                try:
                    parts = ip.split(".")
                    if len(parts) == 4 and parts[2] == self.octet and int(parts[3]) >= self.min_last_octet:
                        rows.append([ip, asset_str, desc_str])
                except: continue
            wb.close()
        except Exception as e: self.error.emit(str(e))
        self.data_loaded.emit(rows, self.target_group)

class PingWorker(QThread):
    ping_done = Signal(str, bool)

    def __init__(self, target, timeout="500"):
        super().__init__()
        self.target = target
        self.timeout = str(timeout)

    def run(self):
        try:
            creation_flags = subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
            proc = subprocess.Popen(
                ["ping", "-n", "1", "-w", self.timeout, self.target],
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, creationflags=creation_flags
            )
            stdout, _ = proc.communicate()
            output = stdout.decode('cp857', errors='ignore')
            online = ("TTL=" in output) or ("Reply from" in output and "Destination host unreachable" not in output)
            self.ping_done.emit(self.target, online)
        except: self.ping_done.emit(self.target, False)

class BulkPingWorker(QThread):
    all_done = Signal(dict)
    one_done = Signal(str, bool)

    def __init__(self, targets: list, timeout="500"):
        super().__init__()
        self.targets = targets
        self.timeout = str(timeout)
        self._results = {}

    def run(self):
        for t in self.targets:
            try:
                creation_flags = subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
                proc = subprocess.Popen(
                    ["ping", "-n", "1", "-w", self.timeout, t],
                    stdout=subprocess.PIPE, stderr=subprocess.PIPE, creationflags=creation_flags
                )
                stdout, _ = proc.communicate()
                out = stdout.decode('cp857', errors='ignore')
                online = ("TTL=" in out) or ("Reply from" in out and "Destination host unreachable" not in out)
            except: online = False
            self._results[t] = online
            self.one_done.emit(t, online)
        self.all_done.emit(self._results)

class PowerWorker(QThread):
    done = Signal(str, bool, str)

    def __init__(self, target, action):
        super().__init__()
        self.target = target; self.action = action

    def run(self):
        arg = "/r" if self.action == "restart" else "/s"
        try:
            creation_flags = subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
            result = subprocess.run(
                ["shutdown", arg, "/f", "/m", f"\\\\{self.target}", "/t", "0"],
                capture_output=True, text=True, creationflags=creation_flags
            )
            success = (result.returncode == 0)
            self.done.emit(self.target, success, "Başarılı" if success else result.stderr.strip())
        except Exception as e:
            self.done.emit(self.target, False, str(e))

class BulkPowerWorker(QThread):
    progress = Signal(str, bool, str)
    all_finished = Signal()

    def __init__(self, targets, action):
        super().__init__()
        self.targets = targets; self.action = action

    def run(self):
        arg = "/r" if self.action == "restart" else "/s"
        for target in self.targets:
            try:
                creation_flags = subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0
                result = subprocess.run(
                    ["shutdown", arg, "/f", "/m", f"\\\\{target}", "/t", "0"],
                    capture_output=True, text=True, creationflags=creation_flags
                )
                success = (result.returncode == 0)
                self.progress.emit(target, success, "Başarılı" if success else result.stderr.strip())
            except Exception as e:
                self.progress.emit(target, False, str(e))
        self.all_finished.emit()

# ─────────────────── ANA PENCERE ───────────────────
class VNCManager(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Flow VNC Controller  ·  by ABŞ")  #biraz da reklam :))
        self.resize(1350, 800)

        self.log_manager = LogManager()
        # FIX: set() → list; add/remove tutarlı kullanımı için
        self.active_ping_threads = []
        self.active_power_threads = set()
        self.connections_today = self._count_today_connections()

        icon_path = get_data_path("img/monitoring.ico")
        if os.path.exists(icon_path): self.setWindowIcon(QIcon(icon_path))

        self.cmrc_enabled = False; self.cmrc_path = ""
        self._load_config()
        self.current_theme = self.config.get("theme", "light")

        self.load_thread = None
        self.current_tree = None
        self.current_group_name = None      
        self.user_folders = {}
        self.ping_status_cache = {}
        self.is_monitoring = False
        self.current_view_is_user = False
        self.bulk_ping_worker = None
        self._prev_offline = set()
        self.favorites = self._load_favorites()
        self.offline_since = {}

        self.auto_monitor_timer = QTimer()
        self.auto_monitor_timer.setInterval(OTOMATIK_IZLEME_DAKIKA * 60 * 1000)
        self.auto_monitor_timer.timeout.connect(self.run_auto_monitor_cycle)

        self.auto_refresh_timer = QTimer()
        if self.config.get("excel_refresh_interval", 0) > 0:
            self.auto_refresh_timer.setInterval(self.config["excel_refresh_interval"] * 60 * 1000)
            self.auto_refresh_timer.timeout.connect(self._auto_refresh_trigger)
            self.auto_refresh_timer.start()

        self._load_user_folders()
        self.init_ui()
        self._setup_tray()

        last_group = self.config.get("last_active_group", "Dashboard")
        self._trigger_group_by_name(last_group)

    # ──  timer ve thread'leri tamamen kapatmak icin eklendi; ──
    def closeEvent(self, event)
        self.auto_monitor_timer.stop()
        self.auto_refresh_timer.stop()
        if self.bulk_ping_worker and self.bulk_ping_worker.isRunning():
            self.bulk_ping_worker.quit()
            self.bulk_ping_worker.wait(2000)
        for w in list(self.active_ping_threads):
            if w.isRunning():
                w.quit(); w.wait(1000)
        self._save_config()
        super().closeEvent(event)

    def show_toast(self, message, notif_type="info"):
        toast = ToastNotification(message, notif_type, self)
        toast.show_at(self)

    def _setup_tray(self):
        icon_path = get_data_path("img/monitoring.ico")
        icon = QIcon(icon_path) if os.path.exists(icon_path) else self.style().standardIcon(QStyle.StandardPixmap.SP_ComputerIcon)
        self.tray = QSystemTrayIcon(icon, self)
        tray_menu = QMenu()
        tray_menu.addAction("Göster", self.show); tray_menu.addAction("Çıkış", QApplication.quit)
        self.tray.setContextMenu(tray_menu); self.tray.show()

    def _count_today_connections(self):
        today = datetime.date.today().strftime("%Y-%m-%d")
        return sum(1 for e in self.log_manager.entries
                   if e.get("action", "").startswith("VNC") and e.get("time", "").startswith(today))

    def init_ui(self):
        central = QWidget(); self.setCentralWidget(central)
        root = QHBoxLayout(central); root.setContentsMargins(0, 0, 0, 0); root.setSpacing(0)

        # ── SOL SIDEBAR ──
        self.sidebar = QWidget(); self.sidebar.setFixedWidth(230); self.sidebar.setObjectName("Sidebar")
        sidebar_lay = QVBoxLayout(self.sidebar); sidebar_lay.setContentsMargins(10, 14, 10, 14); sidebar_lay.setSpacing(4)

        self.logo_label = QLabel(); self.logo_label.setObjectName("LogoLabel")
        logo_f = get_data_path("img/flow_light.png")
        if os.path.exists(logo_f):
            self.logo_label.setPixmap(QPixmap(logo_f).scaledToWidth(120, Qt.SmoothTransformation))
        self.logo_label.setAlignment(Qt.AlignCenter); sidebar_lay.addWidget(self.logo_label)
        sidebar_lay.addSpacing(10)
        sep = QFrame(); sep.setFrameShape(QFrame.HLine); sep.setObjectName("Separator"); sidebar_lay.addWidget(sep)
        sidebar_lay.addSpacing(6)

        self.button_group = QButtonGroup(); self.button_group.setExclusive(True)
        self.dash_btn = SidebarButton("Dashboard", "📊")
        self.dash_btn.clicked.connect(self._show_dashboard); self.button_group.addButton(self.dash_btn); sidebar_lay.addWidget(self.dash_btn)
        self.fav_btn = SidebarButton("Favoriler", "⭐")
        self.fav_btn.clicked.connect(self._show_favorites); self.button_group.addButton(self.fav_btn); sidebar_lay.addWidget(self.fav_btn)
        self.log_btn = SidebarButton("Bağlantı Geçmişi", "📋")
        self.log_btn.clicked.connect(self._show_log_panel); self.button_group.addButton(self.log_btn); sidebar_lay.addWidget(self.log_btn)
        sidebar_lay.addSpacing(6)

        sec_lbl = QLabel("  MONTAJ HATLAR"); sec_lbl.setObjectName("SectionLabel"); sidebar_lay.addWidget(sec_lbl)
        for name in SHEET_MAP.keys():
            btn = DropButton(name, "🖥️")
            btn.clicked.connect(self.create_group_handler(name, is_user=False))
            self.button_group.addButton(btn); sidebar_lay.addWidget(btn)

        sidebar_lay.addSpacing(6)
        sec2_lbl = QLabel("  KLASÖRLER"); sec2_lbl.setObjectName("SectionLabel"); sidebar_lay.addWidget(sec2_lbl)
        self.user_folders_layout = QVBoxLayout(); self.user_folders_layout.setSpacing(4)
        sidebar_lay.addLayout(self.user_folders_layout)
        for uf in self.user_folders.keys(): self._create_and_insert_user_button(uf)
        new_folder_btn = QPushButton("+ Yeni Klasör Ekle"); new_folder_btn.setObjectName("NewFolderBtn")
        new_folder_btn.setCursor(Qt.PointingHandCursor); new_folder_btn.setFixedHeight(34)
        new_folder_btn.clicked.connect(self.on_add_folder); sidebar_lay.addWidget(new_folder_btn)
        sidebar_lay.addStretch()

        self.monitor_btn = QPushButton("🔴  Canlı İzle: KAPALI"); self.monitor_btn.setCheckable(True)
        self.monitor_btn.setFixedHeight(40); self.monitor_btn.setCursor(Qt.PointingHandCursor)
        self.monitor_btn.setObjectName("MonitorBtn"); self.monitor_btn.clicked.connect(self.toggle_monitoring)
        sidebar_lay.addWidget(self.monitor_btn)

        self.bulk_ping_btn = QPushButton("⚡  Toplu Ping"); self.bulk_ping_btn.setFixedHeight(36)
        self.bulk_ping_btn.setCursor(Qt.PointingHandCursor); self.bulk_ping_btn.clicked.connect(self.run_bulk_ping)
        sidebar_lay.addWidget(self.bulk_ping_btn)

        theme_str = "☀️ Aydınlık Tema" if self.current_theme == "dark" else "🌙 Karanlık Tema"
        self.theme_btn = QPushButton(theme_str); self.theme_btn.setFixedHeight(34)
        self.theme_btn.setCursor(Qt.PointingHandCursor); self.theme_btn.clicked.connect(self.toggle_theme)
        sidebar_lay.addWidget(self.theme_btn)

        bottom_row = QHBoxLayout(); bottom_row.addStretch()
        self.bottom_logo_label = QLabel(); self.bottom_logo_label.setObjectName("LogoLabel")
        logo2 = get_data_path("img/duslaci.png")
        if os.path.exists(logo2):
            self.bottom_logo_label.setPixmap(QPixmap(logo2).scaledToWidth(80, Qt.SmoothTransformation))
        bottom_row.addWidget(self.bottom_logo_label); sidebar_lay.addLayout(bottom_row)
        root.addWidget(self.sidebar)

        # ── SAĞ İÇERİK ──
        self.content_area = QWidget(); self.content_area.setObjectName("ContentArea")
        self.content_layout = QVBoxLayout(self.content_area)
        self.content_layout.setContentsMargins(0, 0, 0, 0); root.addWidget(self.content_area, 1)

        self.dashboard_widget = DashboardWidget(self.current_theme)
        self.log_panel = LogPanel(self.log_manager, self, self.current_theme)
        self.favorites_widget = QWidget(); self.connection_widget = QWidget()
        self.stack_widgets = [self.dashboard_widget, self.log_panel, self.favorites_widget, self.connection_widget]
        for w in self.stack_widgets: self.content_layout.addWidget(w); w.hide()

        self.fav_layout = QVBoxLayout(self.favorites_widget)
        self.fav_layout.setContentsMargins(16, 16, 16, 16); self.fav_layout.setSpacing(8)
        self.connection_widget.show()
        self.conn_layout = QVBoxLayout(self.connection_widget)
        self.conn_layout.setContentsMargins(16, 16, 16, 16); self.conn_layout.setSpacing(8)

        self.set_theme(self.current_theme)
        self._setup_shortcuts()

    # ── KLAVYE KISAYOLLARI ──
    def _setup_shortcuts(self):
        QShortcut(QKeySequence("F5"), self, self._shortcut_refresh)
        QShortcut(QKeySequence("Ctrl+F"), self, self._shortcut_search)
        QShortcut(QKeySequence("Return"), self, self._shortcut_enter)

    def _shortcut_refresh(self):
        if self.connection_widget.isVisible() and self.current_group_name and not self.current_view_is_user:
            self.load_sheet_data(self.current_group_name)

    def _shortcut_search(self):
        if hasattr(self, 'search_box') and self.search_box.isVisible():
            self.search_box.setFocus(); self.search_box.selectAll()

    def _shortcut_enter(self):
        if QApplication.activeModalWidget():
            return
        if self.current_tree and self.current_tree.hasFocus() and self.current_tree.currentItem():
            it = self.current_tree.currentItem()
            self.open_vnc(it.text(0), it.text(1))

    def _switch_content(self, widget):
        for w in self.stack_widgets: w.hide()
        widget.show()
        effect = QGraphicsOpacityEffect(widget); widget.setGraphicsEffect(effect)
        anim = QPropertyAnimation(effect, b"opacity"); anim.setDuration(250)
        anim.setStartValue(0.0); anim.setEndValue(1.0); anim.start()
        self._fade_anim = anim

    def _trigger_group_by_name(self, name):
        if name == "Dashboard": self.dash_btn.click()
        elif name == "Favoriler": self.fav_btn.click()
        elif name == "Bağlantı Geçmişi": self.log_btn.click()
        else:
            for b in self.button_group.buttons():
                if hasattr(b, '_base_title') and b._base_title == name:
                    b.click(); return

    def _update_last_group_config(self, group_name):
        self.config["last_active_group"] = group_name; self._save_config()

    def _show_dashboard(self):
        self._uncheck_group_except(self.dash_btn); self.dash_btn.setChecked(True)
        self._update_dashboard_stats(); self._switch_content(self.dashboard_widget)
        self._update_last_group_config("Dashboard")

    def _show_log_panel(self):
        self._uncheck_group_except(self.log_btn); self.log_btn.setChecked(True)
        self.log_panel.refresh(); self._switch_content(self.log_panel)
        self._update_last_group_config("Bağlantı Geçmişi")

    # ── FAVORİLER SEKMESİ ──
    def _load_favorites(self):
        try:
            with open(FAVORITES_FILE, "r", encoding="utf-8") as f: return json.load(f)
        except: return []

    def _save_favorites(self):
        try:
            with open(FAVORITES_FILE, "w", encoding="utf-8") as f:
                json.dump(self.favorites, f, ensure_ascii=False, indent=2)
        except Exception as e: QMessageBox.warning(self, "Hata", f"Favoriler kaydedilemedi:\n{e}")

    def is_favorite(self, ip, asset):
        return any(f["ip"] == ip and f["asset"] == asset for f in self.favorites)

    def toggle_favorite(self, ip, asset, desc):
        if self.is_favorite(ip, asset):
            self.favorites = [f for f in self.favorites if not (f["ip"] == ip and f["asset"] == asset)]
            self.show_toast(f"Favorilerden kaldırıldı: {asset or ip}", "info")
        else:
            self.favorites.append({"ip": ip, "asset": asset, "desc": desc,
                                   "added": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")})
            self.show_toast(f"Favorilere eklendi: {asset or ip} ⭐", "success")
        self._save_favorites()
        if self.favorites_widget.isVisible(): self._show_favorites()

    def _show_favorites(self):
        self._uncheck_group_except(self.fav_btn); self.fav_btn.setChecked(True)
        self._switch_content(self.favorites_widget); self._rebuild_favorites_panel()
        self._update_last_group_config("Favoriler")

    def _rebuild_favorites_panel(self):
        for i in reversed(range(self.fav_layout.count())):
            item = self.fav_layout.takeAt(i)
            if item.widget(): item.widget().deleteLater()

        hdr = QHBoxLayout()
        title = QLabel("⭐  Favoriler"); title.setStyleSheet("font-size:18px; font-weight:bold;")
        hdr.addWidget(title); hdr.addStretch()
        count_lbl = QLabel(f"{len(self.favorites)} kayıt"); count_lbl.setStyleSheet("font-size:11px; color:#6b7280;")
        hdr.addWidget(count_lbl)
        hdr_w = QWidget(); hdr_w.setLayout(hdr); self.fav_layout.addWidget(hdr_w)

        if not self.favorites:
            empty = QLabel("Henüz favori eklenmemiş.\nBir cihaza sağ tıklayıp ⭐ Favoriye Ekle seçin.")
            empty.setAlignment(Qt.AlignCenter); empty.setStyleSheet("color:#9ca3af; font-size:13px; padding:40px;")
            self.fav_layout.addWidget(empty); self.fav_layout.addStretch(); return

        fav_tree = DragTree()
        fav_tree.setHeaderLabels(["IP Adresi", "Asset", "Açıklama", "Eklenme", "Durum", "Süre"])
        fav_tree.setColumnWidth(0, 150); fav_tree.setColumnWidth(1, 160)
        fav_tree.setColumnWidth(2, 250); fav_tree.setColumnWidth(3, 130)
        fav_tree.setColumnWidth(4, 80);  fav_tree.setColumnWidth(5, 80)
        fav_tree.setAlternatingRowColors(False)
        fav_tree.setItemDelegate(StatusDelegate(fav_tree))
        self._apply_tree_style(fav_tree, COLORS[self.current_theme])
        fav_tree.itemDoubleClicked.connect(lambda it, c: self.open_vnc(it.text(0), it.text(1)))
        fav_tree.setContextMenuPolicy(Qt.CustomContextMenu)
        fav_tree.customContextMenuRequested.connect(lambda pos: self._fav_context_menu(fav_tree, pos))

        for fav in self.favorites:
            ip, asset, desc = fav["ip"], fav["asset"], fav["desc"]; added = fav.get("added", "")
            st = self.ping_status_cache.get(ip or asset, "unknown")
            st_text = {"online": "● Online", "offline": "○ Offline", "unknown": "○ ?"}[st]
            dur_text = self._calc_offline_duration(ip or asset) if st == "offline" else ""
            it = QTreeWidgetItem([ip, asset, desc, added, st_text, dur_text])
            it.setData(0, Qt.UserRole, st); fav_tree.addTopLevelItem(it)

        self.fav_layout.addWidget(fav_tree)
        btn_row = QHBoxLayout()
        ping_all_btn = QPushButton("⚡ Hepsini Ping At"); ping_all_btn.setFixedHeight(32)
        ping_all_btn.clicked.connect(lambda: self._ping_favorites(fav_tree))
        btn_row.addWidget(ping_all_btn); btn_row.addStretch()
        btn_w = QWidget(); btn_w.setLayout(btn_row); self.fav_layout.addWidget(btn_w)

    def _fav_context_menu(self, tree, pos):
        item = tree.itemAt(pos)
        if not item: return
        ip, asset = item.text(0).strip(), item.text(1).strip()
        menu = QMenu()
        menu.addAction("🖥️ VNC Bağlan").triggered.connect(lambda: self.open_vnc(ip, asset))
        if self.cmrc_enabled: menu.addAction("🔧 CMRC Bağlan").triggered.connect(lambda: self.open_cmrc(ip, asset))
        menu.addAction("📡 Ping At").triggered.connect(lambda: self.run_ping(item))
        menu.addSeparator()
        menu.addAction("⭐ Favorilerden Kaldır").triggered.connect(lambda: self.toggle_favorite(ip, asset, item.text(2)))
        menu.addAction("📋 IP Kopyala").triggered.connect(lambda: self.copy_to_clipboard(ip))
        menu.exec_(tree.mapToGlobal(pos))

    def _ping_favorites(self, tree):
        targets = {}
        for i in range(tree.topLevelItemCount()):
            it = tree.topLevelItem(i)
            t = it.text(0).strip() or it.text(1).strip()
            if t: targets[t] = it
        if not targets: return
        worker = BulkPingWorker(list(targets.keys()), self.config.get("ping_timeout", 500))
        worker.one_done.connect(lambda tgt, online: self._on_fav_ping(tgt, online, targets, tree))
        worker.start(); self._fav_bulk_worker = worker

    def _on_fav_ping(self, target, online, items, tree):
        # FIX: C++ nesnesi silinmiş olabilir — RuntimeError yakala
        try:
            if tree is None or not tree.isVisible(): return
        except RuntimeError: return

        st = "online" if online else "offline"
        self.ping_status_cache[target] = st
        it = items.get(target)
        try:
            if it is None or it.treeWidget() is None: return
            it.setData(0, Qt.UserRole, st)
            it.setText(4, {"online": "● Online", "offline": "○ Offline"}.get(st, "○ ?"))
            it.setText(5, self._calc_offline_duration(target) if st == "offline" else "")
            tree.viewport().update()
        except RuntimeError: pass

    def _uncheck_group_except(self, keep_btn):
        for b in self.button_group.buttons():
            if b is not keep_btn: b.setChecked(False)

    def _update_dashboard_stats(self):
        on = sum(1 for v in self.ping_status_cache.values() if v == "online")
        off = sum(1 for v in self.ping_status_cache.values() if v == "offline")
        tot = len(self.ping_status_cache)
        self.dashboard_widget.update_stats({"total": tot, "online": on, "offline": off,
                                            "unknown": tot - on - off, "connections_today": self.connections_today})

    # ── CONFIG ──
    def _load_config(self):
        default = {
            "enable_cmrc": True,
            "cmrc_path": r"Sccm Remote Control Tool -1906\\CmRcViewer.exe",
            "ping_timeout": 500,
            "excel_refresh_interval": 0,
            "last_active_group": "Dashboard",
            #"theme": "light",  # kullanıcıların Config dosyasıyla çakışıyor duruma göre açarım bunu
            "column_widths": {"0": 180, "1": 220, "2": 450, "3": 120}
        }
        if os.path.exists(CONFIG_FILE):
            try:
                with open(CONFIG_FILE, "r", encoding="utf-8") as f: self.config = json.load(f)
                for k, v in default.items():
                    if k not in self.config: self.config[k] = v
            except: self.config = default
        else:
            self.config = default; self._save_config()
        self.cmrc_enabled = self.config.get("enable_cmrc", False)
        self.cmrc_path = self.config.get("cmrc_path", "")
        # Geçersiz tema değerini düzelt
        if self.config.get("theme") not in ("light", "dark"):
            self.config["theme"] = "light"
            self._save_config()

    def _save_config(self):
        try:
            with open(CONFIG_FILE, "w", encoding="utf-8") as f: json.dump(self.config, f, indent=4)
        except: pass

    # ── TEMA ──
    def toggle_theme(self):
        new_theme = "dark" if self.current_theme == "light" else "light"
        self.theme_btn.setText("☀️ Aydınlık Tema" if new_theme == "dark" else "🌙 Karanlık Tema")
        self.config["theme"] = new_theme; self._save_config()
        self.set_theme(new_theme)

    def set_theme(self, mode="light"):
        self.current_theme = mode; C = COLORS[mode]

        logo_file = "flow_dark.png" if mode == "dark" else "flow_light.png"
        lf = get_data_path(f"img/{logo_file}")
        if os.path.exists(lf):
            self.logo_label.setPixmap(QPixmap(lf).scaledToWidth(120, Qt.SmoothTransformation))

        self.sidebar.setStyleSheet(
            f"QWidget#Sidebar {{ background: {C['bg_sidebar']}; border-right: 1px solid {C['border']}; }}"
            f"QLabel#SectionLabel {{ color: {C['text_dim']}; font-size: 10px; font-weight: bold; letter-spacing: 1px; padding: 4px 0 2px 4px; }}"
            f"QFrame#Separator {{ color: {C['border']}; }}"
            f"QPushButton {{ background: transparent; color: {C['text']}; border: none; border-radius: 8px; text-align: left; padding-left: 8px; font-size: 13px; }}"
            f"QPushButton:hover {{ background: {C['btn_hover']}; }}"
            f"QPushButton:checked {{ background: {C['accent']}; color: white; font-weight: bold; }}"
            f"QPushButton#NewFolderBtn {{ background: {C['btn_bg']}; color: {C['text_dim']}; border: 1px dashed {C['border']}; border-radius: 8px; font-size: 12px; text-align: center; }}"
            f"QPushButton#NewFolderBtn:hover {{ background: {C['btn_hover']}; color: {C['text']}; }}"
            f"QPushButton#MonitorBtn {{ text-align: left; font-weight: bold; border-radius: 8px; }}"
            f"QLabel#LogoLabel {{ background-color: {C['logo_bg']}; border-radius: 6px; padding: 4px; }}"
        )

        self.centralWidget().setStyleSheet(
            f"* {{ color: {C['text']}; }}"
            f"QWidget#ContentArea, QWidget#ContentArea QWidget {{ background: {C['bg_main']}; }}"
            f"QLabel {{ color: {C['text']}; background: transparent; }}"
            f"QLabel#StatusTextLabel {{ color: {C['text']}; font-size: 11px; background: transparent; }}"
            f"QLineEdit {{ background: {C['bg_card']}; color: {C['text']}; border: 1px solid {C['border']}; border-radius: 8px; padding: 6px 10px; font-size: 13px; }}"
            f"QLineEdit:focus {{ border: 1px solid {C['accent']}; }}"
            f"QPushButton {{ background: {C['btn_bg']}; color: {C['text']}; border: 1px solid {C['border']}; border-radius: 8px; padding: 5px 14px; font-size: 13px; }}"
            f"QPushButton:hover {{ background: {C['btn_hover']}; }}"
            f"QScrollArea {{ border: none; background: transparent; }}"
            f"QProgressBar {{ border-radius: 6px; background: {C['bg_card']}; }}"
            f"QProgressBar::chunk {{ border-radius: 6px; background: qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 {C['online']},stop:1 {C['accent']}); }}"
            f"QScrollBar:vertical {{ background: {C['bg_sidebar']}; width: 6px; border-radius: 3px; }}"
            f"QScrollBar::handle:vertical {{ background: {C['border']}; border-radius: 3px; min-height: 30px; }}"
        )

        for lbl in self.centralWidget().findChildren(QLabel):
            ss = lbl.styleSheet()
            if lbl.parent() and lbl.parent().objectName() == "Sidebar": continue
            if "online" in ss or "offline" in ss: continue
            # Güvenli split
            size   = ss.split('font-size:')[1].split(';')[0].strip() if 'font-size:' in ss else '13px'
            weight = 'bold' if 'bold' in ss else 'normal'
            lbl.setStyleSheet(f"color: {C['text']}; background: transparent; font-size:{size}; font-weight:{weight};")

        if self.current_tree: self._apply_tree_style(self.current_tree, C)
        self.dashboard_widget.apply_theme(mode)
        self.log_panel.apply_theme(mode)
        self._update_monitor_btn_style()
        if self.connection_widget.isVisible(): self.update_status_counts()

    def _apply_tree_style(self, tree, C):
        tree.setStyleSheet(
            f"QTreeWidget {{ background: {C['bg_main']}; color: {C['text']}; border: none; outline: none; }}"
            f"QTreeWidget::item {{ min-height: 32px; padding: 4px 6px; border-bottom: 1px solid {C['border']}; border-radius: 0px; }}"
            f"QTreeWidget::item:hover {{ background: {C['hover']}; color: {C['text']}; }}"
            f"QTreeWidget::item:selected {{ background: #1a56db; color: #ffffff; }}"
            f"QTreeWidget::item:selected:hover {{ background: #1e4fc2; color: #ffffff; }}"
            f"QTreeWidget::item:alternate {{ background: {C['bg_card']}; }}"
            f"QHeaderView::section {{ background: {C['header_bg']}; color: {C['text_dim']}; padding: 8px 6px; font-weight: bold; font-size: 11px; border: none; border-bottom: 2px solid {C['accent']}; letter-spacing: 0.5px; }}"
        )

        from PySide6.QtGui import QPalette
        pal = tree.palette()
        pal.setColor(QPalette.ColorRole.Base, QColor(C['bg_main']))
        tree.setPalette(pal)

    # ── USER KLASÖRLERI ──
    def _load_user_folders(self):
        if os.path.exists(USER_FOLDER_FILE):
            try:
                with open(USER_FOLDER_FILE, "r", encoding="utf-8") as f: self.user_folders = json.load(f)
            except: self.user_folders = {}

    def _save_user_folders(self):
        try:
            with open(USER_FOLDER_FILE, "w", encoding="utf-8") as f:
                json.dump(self.user_folders, f, ensure_ascii=False, indent=2)
        except Exception as e: QMessageBox.warning(self, "Hata", f"Kaydedilemedi:\n{e}")

    def _create_and_insert_user_button(self, name):
        btn = DropButton(name, "📁", drop_callback=self._on_drop_to_folder)
        btn.setCheckable(True); btn._base_title = name
        btn.clicked.connect(self.create_group_handler(name, is_user=True))
        btn.setContextMenuPolicy(Qt.CustomContextMenu)
        btn.customContextMenuRequested.connect(lambda pos, b=btn: self._show_left_context(b, pos))
        self.user_folders_layout.addWidget(btn); self.button_group.addButton(btn)

    def _show_left_context(self, btn, pos):
        menu = QMenu()
        ren = menu.addAction("✏️ Yeniden Adlandır"); dele = menu.addAction("🗑️ Sil")
        act = menu.exec_(btn.mapToGlobal(pos))
        if act == dele and QMessageBox.question(self, "Onay", f"'{btn._base_title}' silinsin mi?",
                                                QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
            name = btn._base_title; btn.setParent(None)
            self.user_folders.pop(name, None); self._save_user_folders()
        elif act == ren:
            old = btn._base_title
            new, ok = QInputDialog.getText(self, "Düzenle", "Yeni ad:", text=old)
            if ok and new.strip() and new != old:
                self.user_folders[new] = self.user_folders.pop(old, []); self._save_user_folders()
                btn.setText(f"  📁  {new}"); btn._base_title = new

    def on_add_folder(self):
        name, ok = QInputDialog.getText(self, "Yeni Klasör", "Klasör adı:")
        if ok and name.strip() and name not in self.user_folders:
            self.user_folders[name] = []; self._save_user_folders()
            self._create_and_insert_user_button(name)

    # ── GROUP HANDLER ──
    def create_group_handler(self, name, is_user=False):
        def h():
            self.current_group_name = name     # FIX: aktif grubu kaydet
            self.current_view_is_user = is_user
            self._uncheck_group_except(None)
            for b in self.button_group.buttons():
                if hasattr(b, '_base_title'): b.setChecked(b._base_title == name)
            self._switch_content(self.connection_widget)
            self.show_group_connections(name, is_user)
            self._update_last_group_config(name)
        return h

    # ── BAĞLANTI LİSTESİ ──
    def show_group_connections(self, group_name, is_user=False):
        for i in reversed(range(self.conn_layout.count())):
            item = self.conn_layout.takeAt(i)
            if item.widget(): item.widget().deleteLater()

        tbar = QWidget(); tlay = QHBoxLayout(tbar); tlay.setContentsMargins(0, 0, 0, 0)
        self.current_title_label = QLabel(group_name); self.current_title_label.setObjectName("GroupTitleLabel")
        self.current_title_label.setStyleSheet(
            f"font-size:16px; font-weight:bold; color:{COLORS[self.current_theme]['text']}; background:transparent;")
        tlay.addWidget(self.current_title_label); tlay.addStretch()

        if not is_user:
            rb = QPushButton("🔄 Yenile (F5)"); rb.setFixedHeight(30)
            rb.clicked.connect(lambda: self.load_sheet_data(group_name)); tlay.addWidget(rb)

        csv_btn = QPushButton("💾 CSV Çıkar"); csv_btn.setFixedHeight(30)
        csv_btn.clicked.connect(self._export_current_tree_csv); tlay.addWidget(csv_btn)

        pwr_btn = QPushButton("⚡ Toplu Güç"); pwr_btn.setFixedHeight(30)
        pwr_btn.clicked.connect(self.show_bulk_power_dialog); tlay.addWidget(pwr_btn)

        qb = QPushButton("🖥️ Hızlı Bağlan"); qb.setFixedHeight(30)
        qb.clicked.connect(self._show_quick_connect_dialog); tlay.addWidget(qb)
        self.conn_layout.addWidget(tbar)

        status_bar_w = QWidget(); slay = QHBoxLayout(status_bar_w); slay.setContentsMargins(0, 0, 0, 4)
        self.status_bar = QProgressBar(); self.status_bar.setTextVisible(False); self.status_bar.setFixedHeight(6)
        self.status_text_label = QLabel("Bekleniyor..."); self.status_text_label.setObjectName("StatusTextLabel")
        slay.addWidget(self.status_bar, 1); slay.addWidget(self.status_text_label)
        self.conn_layout.addWidget(status_bar_w)

        self.search_box = QLineEdit()
        self.search_box.setPlaceholderText("🔍  Ara (IP, Asset, Açıklama)... (Ctrl+F)")
        self.conn_layout.addWidget(self.search_box)

        self.current_tree = DragTree()
        self.current_tree.setHeaderLabels(["IP Adresi", "Asset", "Açıklama", "Durum Süresi"])
        self.current_tree.setContextMenuPolicy(Qt.CustomContextMenu)
        self.current_tree.itemDoubleClicked.connect(lambda it, c: self.open_vnc(it.text(0), it.text(1)))
        self.current_tree.customContextMenuRequested.connect(self.show_context_menu)

        # Sütun genişlikleri config'den yükle
        c_widths = self.config.get("column_widths", {"0": 180, "1": 220, "2": 450, "3": 120})
        self.current_tree.setColumnWidth(0, int(c_widths.get("0", 180)))
        self.current_tree.setColumnWidth(1, int(c_widths.get("1", 220)))
        self.current_tree.setColumnWidth(2, int(c_widths.get("2", 450)))
        self.current_tree.setColumnWidth(3, int(c_widths.get("3", 120)))
        self.current_tree.header().sectionResized.connect(self._on_col_resized)

        self.current_tree.setAlternatingRowColors(False)
        self._tree_delegate = StatusDelegate(self.current_tree)
        self.current_tree.setItemDelegate(self._tree_delegate)
        self.conn_layout.addWidget(self.current_tree)

        self.search_box.textChanged.connect(lambda t: self.filter_tree(self.current_tree, t))
        self._apply_tree_style(self.current_tree, COLORS[self.current_theme])

        if is_user:
            ents = sorted(self.user_folders.get(group_name, []), key=lambda x: x[0])
            for r in ents:
                r_ext = r + [""] if len(r) == 3 else r
                it = QTreeWidgetItem(r_ext); tar = r[0] or r[1]
                st = self.ping_status_cache.get(tar, "unknown")
                it.setData(0, Qt.UserRole, st)
                if st == "offline": it.setText(3, self._calc_offline_duration(tar))
                self.current_tree.addTopLevelItem(it)
            self.update_status_counts(); return

        self.load_sheet_data(group_name)

    def _on_col_resized(self, index, oldSize, newSize):
        if "column_widths" not in self.config: self.config["column_widths"] = {}
        self.config["column_widths"][str(index)] = newSize
        self._save_config()

    def _export_current_tree_csv(self):
        if not self.current_tree: return
        path, _ = QFileDialog.getSaveFileName(
            self, "Listeyi Dışa Aktar",
            f"{self.current_group_name or 'Cihazlar'}_Liste.csv", "CSV Files (*.csv)")
        if not path: return
        try:
            with open(path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["IP Adresi", "Asset", "Açıklama", "Durum", "Süre"])
                for i in range(self.current_tree.topLevelItemCount()):
                    it = self.current_tree.topLevelItem(i)
                    st = it.data(0, Qt.UserRole)
                    writer.writerow([it.text(0), it.text(1), it.text(2), st, it.text(3)])
            self.show_toast("Liste dışa aktarıldı.", "success")
        except Exception as e:
            self.show_toast(f"Dışa aktarma hatası: {e}", "error")

    # ── CONTEXT MENÜ ──
    def show_context_menu(self, pos):
        item = self.current_tree.itemAt(pos)
        menu = QMenu()
        if item:
            ip, asst = item.text(0).strip(), item.text(1).strip()
            if ip or asst:
                menu.addAction("📡 Ping At").triggered.connect(lambda: self.run_ping(item))
                menu.addAction("🖥️ VNC Bağlan  [Enter]").triggered.connect(lambda: self.open_vnc(ip, asst))
                fav_label = "⭐ Favorilerden Kaldır" if self.is_favorite(ip, asst) else "⭐ Favoriye Ekle"
                menu.addAction(fav_label).triggered.connect(
                    lambda: self.toggle_favorite(ip, asst, item.text(2) if item.columnCount() > 2 else ""))
                if self.cmrc_enabled:
                    menu.addAction("🔧 CMRC Bağlan").triggered.connect(lambda: self.open_cmrc(ip, asst))
                menu.addSeparator()
                if self.current_view_is_user:
                    menu.addAction("✏️ Düzenle").triggered.connect(lambda: self.edit_item(item))
                    menu.addAction("🗑️ Sil").triggered.connect(lambda: self.delete_item(item))
                copy = menu.addMenu("📋 Kopyala")
                if ip: copy.addAction("IP Kopyala").triggered.connect(lambda: self.copy_to_clipboard(ip))
                if asst: copy.addAction("Asset Kopyala").triggered.connect(lambda: self.copy_to_clipboard(asst))
                pwr = menu.addMenu("⚡ Güç İşlemi")
                pwr.addAction("🔄 Yeniden Başlat").triggered.connect(lambda: self.remote_power_action(ip, asst, "restart"))
                pwr.addAction("🛑 Kapat").triggered.connect(lambda: self.remote_power_action(ip, asst, "shutdown"))
        elif self.current_view_is_user:
            menu.addAction("➕ Ekle").triggered.connect(self.add_connection)
        menu.exec_(self.current_tree.mapToGlobal(pos))

    # ── GÜÇ İŞLEMLERİ ──
    def remote_power_action(self, ip, asst, act):
        t = ip or asst
        if t and QMessageBox.warning(self, "Onay", f"{t} cihazına komut gönderilecek. Emin misiniz?",
                                     QMessageBox.Yes | QMessageBox.No) == QMessageBox.Yes:
            worker = PowerWorker(t, act)
            self.active_power_threads.add(worker)
            worker.done.connect(self._on_power_done)
            worker.finished.connect(lambda w=worker: self.active_power_threads.discard(w))
            worker.start()
            self.show_toast(f"Komut gönderiliyor: {t}", "info")

    def show_bulk_power_dialog(self):
        if not self.current_tree: return
        targets = []
        for i in range(self.current_tree.topLevelItemCount()):
            it = self.current_tree.topLevelItem(i)
            if it.data(0, Qt.UserRole) == "online":
                t = it.text(0).strip() or it.text(1).strip()
                if t: targets.append(t)
        if not targets:
            self.show_toast("Online durumda cihaz bulunamadı.", "warning"); return

        dia = QDialog(self); dia.setWindowTitle("Toplu Güç İşlemi"); dia.setFixedSize(350, 130)
        lay = QVBoxLayout(dia)
        lay.addWidget(QLabel(f"Toplam {len(targets)} adet ONLINE cihaza toplu komut gönderilecek.\nLütfen işlemi seçin:"))
        btn_lay = QHBoxLayout()
        r_btn = QPushButton("🔄 Yeniden Başlat"); s_btn = QPushButton("🛑 Kapat")
        btn_lay.addWidget(r_btn); btn_lay.addWidget(s_btn); lay.addLayout(btn_lay)
        r_btn.clicked.connect(lambda: (self._run_bulk_power(targets, "restart"), dia.accept()))
        s_btn.clicked.connect(lambda: (self._run_bulk_power(targets, "shutdown"), dia.accept()))
        dia.exec()

    def _run_bulk_power(self, targets, action):
        self.show_toast(f"Toplu komut başlatıldı ({len(targets)} cihaz).", "info")
        self.bulk_pwr_worker = BulkPowerWorker(targets, action)
        self.bulk_pwr_worker.progress.connect(self._on_power_done)
        self.bulk_pwr_worker.all_finished.connect(lambda: self.show_toast("Toplu güç işlemi tamamlandı.", "success"))
        self.bulk_pwr_worker.start()

    def _on_power_done(self, target, success, msg):
        if success:
            self.show_toast(f"✅ Başarılı: {target}", "success")
            self.log_manager.add("Güç İşlemi", target, "Başarılı", "ok")
        else:
            self.show_toast(f"❌ Hata ({target}): {msg}", "error")
            self.log_manager.add("Güç Hatası", target, msg, "error")

    # ── TOPLU PİNG ──
    def run_bulk_ping(self):
        # Önceki pingleme bitmeden yeni başlatmıyoruz
        if self.bulk_ping_worker and self.bulk_ping_worker.isRunning():
            self.show_toast("Mevcut tarama işlemi henüz bitmedi.", "warning"); return
        if not self.current_tree:
            self.show_toast("Önce bir grup seçin.", "warning"); return

        targets = {}
        for i in range(self.current_tree.topLevelItemCount()):
            it = self.current_tree.topLevelItem(i)
            t = it.text(0).strip() or it.text(1).strip()
            if t: targets[t] = it
        if not targets:
            self.show_toast("Ping atılacak cihaz yok.", "warning"); return

        self.bulk_ping_btn.setText("⏳  Taranıyor..."); self.bulk_ping_btn.setEnabled(False)
        self.bulk_ping_worker = BulkPingWorker(list(targets.keys()), self.config.get("ping_timeout", 500))
        self.bulk_ping_worker.one_done.connect(lambda tgt, online: self._on_bulk_one(tgt, online, targets))
        self.bulk_ping_worker.all_done.connect(self._on_bulk_done)
        self.bulk_ping_worker.start()

    def _on_bulk_one(self, target, online, items):
        if not self.current_tree: return
        st = "online" if online else "offline"
        was_online = self.ping_status_cache.get(target)
        self.ping_status_cache[target] = st

        if st == "offline":
            if target not in self.offline_since:
                self.offline_since[target] = datetime.datetime.now()
        elif target in self.offline_since:
            self.offline_since.pop(target, None)

        it = items.get(target)
        try:
            if it is None or it.treeWidget() is None: return
            it.setData(0, Qt.UserRole, st)
            it.setText(3, self._calc_offline_duration(target) if st == "offline" else "")
        except RuntimeError: return

        try: self.update_status_counts()
        except RuntimeError: pass

    def _on_bulk_done(self, results):
        self.bulk_ping_btn.setText("⚡  Toplu Ping"); self.bulk_ping_btn.setEnabled(True)
        on = sum(1 for v in results.values() if v); off = len(results) - on
        self.show_toast(f"Tarama tamamlandı: {on} online, {off} offline", "success")
        offline_set = {t for t, v in results.items() if not v}
        newly_down = offline_set - self._prev_offline
        if newly_down: self._notify_offline(newly_down)
        self._prev_offline = offline_set

    def _calc_offline_duration(self, target):
        if target not in self.offline_since: return ""
        dur = datetime.datetime.now() - self.offline_since[target]
        mins = int(dur.total_seconds() // 60)
        secs = int(dur.total_seconds() % 60)
        if mins == 0: return f"{secs}sn"
        if mins < 60: return f"{mins}dk {secs}sn"
        return f"{mins // 60}sa {mins % 60}dk"

    def _notify_offline(self, targets):
        names = ', '.join(list(targets)[:3]) + ('...' if len(targets) > 3 else '')
        msg = f"⚠️ {names} çevrimdışı oldu!"
        if hasattr(self, "tray") and self.tray.isSystemTrayAvailable():
            self.tray.showMessage("Flow VNC — Cihaz Uyarısı", msg, QSystemTrayIcon.Warning, 5000)
        self.show_toast(msg, "error")
        for t in targets:
            if t not in self.offline_since:
                self.offline_since[t] = datetime.datetime.now()
            self.log_manager.add("Offline Uyarı", t, "Online'dan offline'a düştü", "error")
        if HAS_WINSOUND:
            try: winsound.MessageBeep(winsound.MB_ICONEXCLAMATION)
            except: pass

    # ── İZLEME ──
    def _update_monitor_btn_style(self):
        C = COLORS[self.current_theme]
        color = C["online"] if self.is_monitoring else C["offline"]
        self.monitor_btn.setText(
            f"🟢  Canlı İzle: AÇIK ({OTOMATIK_IZLEME_DAKIKA}dk)" if self.is_monitoring else "🔴  Canlı İzle: KAPALI")
        self.monitor_btn.setStyleSheet(
            f"text-align:left; padding:8px 12px; font-weight:bold; color:{color};"
            f"background:transparent; border:{'1px solid '+color if self.is_monitoring else 'none'}; border-radius:8px;")

    def toggle_monitoring(self):
        self.is_monitoring = self.monitor_btn.isChecked()
        if self.is_monitoring:
            self.auto_monitor_timer.start(); self.run_auto_monitor_cycle()
        else: self.auto_monitor_timer.stop()
        self._update_monitor_btn_style()

    def run_auto_monitor_cycle(self):
        if self.bulk_ping_worker and self.bulk_ping_worker.isRunning(): return
        if not self.current_tree: return
        targets = {}
        for i in range(self.current_tree.topLevelItemCount()):
            it = self.current_tree.topLevelItem(i)
            tar = it.text(0).strip() or it.text(1).strip()
            if not it.isHidden() and tar: targets[tar] = it
        if not targets: return
        self.bulk_ping_worker = BulkPingWorker(list(targets.keys()), self.config.get("ping_timeout", 500))
        self.bulk_ping_worker.one_done.connect(lambda tgt, online: self._on_bulk_one(tgt, online, targets))
        self.bulk_ping_worker.all_done.connect(self._on_monitor_done)
        self.bulk_ping_worker.start()

    def _on_monitor_done(self, results):
        offline_set = {t for t, v in results.items() if not v}
        newly_down = offline_set - self._prev_offline
        if newly_down: self._notify_offline(newly_down)
        # Tekrar online olanlar toast bildirim gönder
        newly_up = self._prev_offline - offline_set
        for t in newly_up:
            if t in self.offline_since:
                dur = datetime.datetime.now() - self.offline_since.pop(t)
                self.show_toast(f"✅ {t} tekrar online! ({self._calc_duration(dur)} offline kaldı)", "success")
                self.log_manager.add("Tekrar Online", t, f"{self._calc_duration(dur)} offline kaldı", "ok")
        self._prev_offline = offline_set

    def _calc_duration(self, delta):
        total = int(delta.total_seconds())
        if total < 60: return f"{total}sn"
        if total < 3600: return f"{total//60}dk {total%60}sn"
        return f"{total//3600}sa {(total%3600)//60}dk"

    def _auto_refresh_trigger(self):
        if self.connection_widget.isVisible() and self.current_group_name and not self.current_view_is_user:
            self.load_sheet_data(self.current_group_name)

    # ── HIZLI BAĞLANTI ──
    def _show_quick_connect_dialog(self):
        dia = QDialog(self); dia.setWindowTitle("Hızlı Bağlantı"); dia.setFixedSize(320, 130)
        lay = QVBoxLayout(dia)
        inp = QLineEdit(); inp.setPlaceholderText("IP adresi veya Asset...")
        lay.addWidget(QLabel("Hedef:")); lay.addWidget(inp)
        bl = QHBoxLayout(); lay.addLayout(bl)
        vb = QPushButton("🖥️ UltraVNC"); bl.addWidget(vb)
        if self.cmrc_enabled:
            cb = QPushButton("🔧 CmRc"); bl.addWidget(cb)
            cb.clicked.connect(lambda: (self.open_cmrc("", inp.text().strip()), dia.accept()))
        vb.clicked.connect(lambda: (self.open_vnc(inp.text().strip()), dia.accept()))
        dia.exec()

    # ── EXCEL YÜKLEME ──
    def load_sheet_data(self, gname):
        if self.load_thread and self.load_thread.isRunning():
            self.show_toast("Veriler zaten yükleniyor...", "warning"); return

        load_lbl = QLabel("⏳  Yükleniyor..."); load_lbl.setAlignment(Qt.AlignCenter)
        load_lbl.setStyleSheet("font-size:14px; color: #6b7280; padding: 40px;")
        self.conn_layout.addWidget(load_lbl); QApplication.processEvents()

        min_last = 1 if OCTET_MAP.get(gname, "") == "39" else 135
        self.load_thread = LoadSheetWorker(SHEET_MAP.get(gname, ""), OCTET_MAP.get(gname, ""), gname, min_last)
        self.load_thread.data_loaded.connect(
            lambda rows, grp: self._on_rows_loaded(rows, load_lbl, grp))
        self.load_thread.error.connect(lambda msg: self.show_toast(f"Excel Hatası: {msg}", "error"))
        self.load_thread.start()

    def _on_rows_loaded(self, rows, label, target_group):
        if target_group != self.current_group_name:
            try:
                if label and label.parent(): label.deleteLater()
            except RuntimeError: pass
            return

        try:
            if label and label.parent(): label.deleteLater()
        except RuntimeError: pass

        if self.current_tree:
            self.current_tree.clear()
            for r in rows:
                r_ext = r + [""] if len(r) == 3 else r
                it = QTreeWidgetItem(r_ext); tar = r[0] or r[1]
                st = self.ping_status_cache.get(tar, "unknown")
                it.setData(0, Qt.UserRole, st)
                if st == "offline": it.setText(3, self._calc_offline_duration(tar))
                self.current_tree.addTopLevelItem(it)
            self.update_status_counts()

    # ── STATUS SAYAÇ ──
    def update_status_counts(self):
        on, off, tot = 0, 0, 0
        C = COLORS[self.current_theme]
        if self.current_tree:
            for i in range(self.current_tree.topLevelItemCount()):
                it = self.current_tree.topLevelItem(i)
                if it.isHidden() or not (it.text(0) or it.text(1)): continue
                st = it.data(0, Qt.UserRole); tot += 1
                if st == "online": on += 1
                elif st == "offline": off += 1
        if hasattr(self, "status_bar") and tot > 0: self.status_bar.setValue(int((on / tot) * 100))
        if hasattr(self, "status_text_label"):
            self.status_text_label.setStyleSheet(f"color: {C['text']}; font-size:11px; background: transparent;")
            self.status_text_label.setText(
                f"<span style='color:{C['text']}'>Toplam: {tot} | "
                f"<span style='color:{C['online']}'>● {on} online</span> | "
                f"<span style='color:{C['offline']}'>● {off} offline</span></span>"
            )

    # ── FİLTRELE ──
    def filter_tree(self, tree, text):
        if not tree: return
        for i in range(tree.topLevelItemCount()):
            it = tree.topLevelItem(i)
            it.setHidden(not any(text.lower() in it.text(c).lower() for c in range(3)))
        self.update_status_counts()

    # ── DÜZENLEME / SİLME ──
    def copy_to_clipboard(self, t): QGuiApplication.clipboard().setText(t)

    def edit_item(self, item):
        old_ip, old_as, old_ds = item.text(0), item.text(1), item.text(2)
        nip, ok1 = QInputDialog.getText(self, "Düzenle", "IP Adresi:", text=old_ip)
        if not ok1: return
        nas, ok2 = QInputDialog.getText(self, "Düzenle", "Asset:", text=old_as)
        if not ok2: return
        nds, ok3 = QInputDialog.getText(self, "Düzenle", "Açıklama:", text=old_ds)
        if not ok3: return
        item.setText(0, nip); item.setText(1, nas); item.setText(2, nds)
        p = next((b._base_title for b in self.button_group.buttons() if b.isChecked() and hasattr(b, '_base_title')), None)
        if p in self.user_folders:
            for i, e in enumerate(self.user_folders[p]):
                if e[0] == old_ip and e[1] == old_as: self.user_folders[p][i] = [nip, nas, nds]; break
            self._save_user_folders()

    def delete_item(self, item):
        p = next((b._base_title for b in self.button_group.buttons() if b.isChecked() and hasattr(b, '_base_title')), None)
        self.current_tree.takeTopLevelItem(self.current_tree.indexOfTopLevelItem(item))
        if p in self.user_folders:
            self.user_folders[p] = [e for e in self.user_folders[p] if not (e[0] == item.text(0) and e[1] == item.text(1))]
            self._save_user_folders()
        self.update_status_counts()

    def _on_drop_to_folder(self, folder, ip, asset, desc):
        if not (ip or asset): return
        arr = self.user_folders.get(folder, [])
        if not any(e[0] == ip and e[1] == asset for e in arr):
            arr.append([ip, asset, desc]); self.user_folders[folder] = arr; self._save_user_folders()
            active_btn = next((b for b in self.button_group.buttons() if b.isChecked()), None)
            if active_btn and hasattr(active_btn, '_base_title') and active_btn._base_title == folder:
                self.show_group_connections(folder, True)

    # ── TEK PİNG ──
    def run_ping(self, it):
        t = it.text(0).strip() or it.text(1).strip()
        if not t: return
        worker = PingWorker(t, self.config.get("ping_timeout", 500))
        self.active_ping_threads.append(worker)
        worker.ping_done.connect(lambda target, online, item=it: self._on_ping_done(item, online, target))
        worker.finished.connect(self._cleanup_ping_thread)
        worker.start()

    def _cleanup_ping_thread(self):
        worker = self.sender()
        try: self.active_ping_threads.remove(worker)
        except ValueError: pass

    def _on_ping_done(self, it, online, target):
        st = "online" if online else "offline"
        was_online = self.ping_status_cache.get(target)
        self.ping_status_cache[target] = st

        try:
            if it is None or it.treeWidget() is None: return
            it.setData(0, Qt.UserRole, st)
            if st == "offline":
                if target not in self.offline_since:
                    self.offline_since[target] = datetime.datetime.now()
                it.setText(3, self._calc_offline_duration(target))
            else:
                if target in self.offline_since: self.offline_since.pop(target, None)
                it.setText(3, "")
            self.update_status_counts()
        except RuntimeError: pass

        if was_online == "online" and not online:
            self._notify_offline({target})
        elif online and was_online == "offline":
            if target in self.offline_since:
                dur = datetime.datetime.now() - self.offline_since.pop(target)
                self.show_toast(f"✅ {target} tekrar online! ({self._calc_duration(dur)} offline kaldı)", "success")
                self.log_manager.add("Tekrar Online", target, f"{self._calc_duration(dur)} offline kaldı", "ok")

    # ── EKLE ──
    def add_connection(self):
        ip, ok1 = QInputDialog.getText(self, "Ekle", "IP Adresi:")
        if not ok1: return
        asst, ok2 = QInputDialog.getText(self, "Ekle", "Hostname:")
        if not ok2: return
        desc, ok3 = QInputDialog.getText(self, "Ekle", "Not:")
        if not ok3: return
        self.current_tree.addTopLevelItem(QTreeWidgetItem([ip, asst, desc, ""]))
        p = next((b._base_title for b in self.button_group.buttons() if b.isChecked() and hasattr(b, '_base_title')), None)
        if p in self.user_folders: self.user_folders[p].append([ip, asst, desc]); self._save_user_folders()

    # ── VNC / CMRC ──
    def open_vnc(self, ip, asst=""):
        t = ip or asst
        if not t: return
        ps = [r"C:\Program Files\uvnc bvba\UltraVNC\vncviewer.exe", r"C:\Program Files\UltraVNC\vncviewer.exe"]
        v = next((p for p in ps if os.path.exists(p)), None)
        if v:
            subprocess.Popen([v, t, "-password", VNC_PASSWORD])
            self.log_manager.add("VNC Bağlantı", t, asst, "ok")
            self.connections_today += 1; self.show_toast(f"VNC bağlanıyor: {t}", "success")
        else:
            self.log_manager.add("VNC Bağlantı", t, "Viewer bulunamadı", "error")
            QMessageBox.critical(self, "Hata", "UltraVNC Viewer bulunamadı!")

    def open_cmrc(self, ip, asst):
        if self.cmrc_path and os.path.exists(self.cmrc_path):
            subprocess.Popen([self.cmrc_path, asst or ip], creationflags=subprocess.CREATE_NO_WINDOW)
            self.log_manager.add("CMRC Bağlantı", asst or ip, "", "ok")
            self.show_toast(f"CMRC bağlanıyor: {asst or ip}", "info")
        else: QMessageBox.warning(self, "Hata", f"CMRC yolu geçersiz:\n{self.cmrc_path}")

# ─────────────────── GİRİŞ ───────────────────
if __name__ == "__main__":
    import multiprocessing
    multiprocessing.freeze_support()

    app = QApplication(sys.argv)
    app.setStyle("Fusion")

    # ── Splash Screen ──
    W, H = 480, 280
    splash_pix = QPixmap(W, H)
    splash_pix.fill(QColor("#f0f2f8"))

    painter = QPainter(splash_pix)
    painter.setRenderHint(QPainter.Antialiasing)

    # Üst accent şerit
    painter.fillRect(0, 0, W, 5, QColor("#3b7cf8"))

    # Ana başlık
    painter.setPen(QColor("#1a1d2e"))
    font = painter.font()
    font.setPointSize(22)
    font.setBold(True)
    painter.setFont(font)
    painter.drawText(QRect(0, 30, W, 50), Qt.AlignHCenter | Qt.AlignVCenter, "Flow VNC Controller")

    # Alt başlık
    font.setPointSize(11)
    font.setBold(False)
    painter.setFont(font)
    painter.setPen(QColor("#6b7280"))
    painter.drawText(QRect(0, 80, W, 30), Qt.AlignHCenter | Qt.AlignVCenter, "by ABŞ")

    # İnce ayırıcı çizgi
    painter.setPen(QColor("#dde1ec"))
    painter.drawLine(40, 125, W - 40, 125)

    # duslaci.png logosunu ortala ve ekle
    duslaci_path = get_data_path("img/duslaci.png")
    if os.path.exists(duslaci_path):
        logo = QPixmap(duslaci_path).scaledToWidth(110, Qt.SmoothTransformation)
        lx = (W - logo.width()) // 2
        painter.drawPixmap(lx, 140, logo)

    # "Yükleniyor..." yazısı
    font.setPointSize(10)
    painter.setFont(font)
    painter.setPen(QColor("#9ca3af"))
    painter.drawText(QRect(0, H - 36, W, 24), Qt.AlignHCenter | Qt.AlignVCenter, "Yükleniyor...")

    # Alt accent şerit
    painter.fillRect(0, H - 4, W, 4, QColor("#3b7cf8"))

    painter.end()

    splash = QSplashScreen(splash_pix, Qt.WindowStaysOnTopHint)
    splash.show()
    app.processEvents()

    # Pencereyi oluştur
    window = VNCManager()
    window.show()
    splash.finish(window)   # Program hazır oldugunda splash kapanır

    sys.exit(app.exec())
